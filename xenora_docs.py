from __future__ import annotations

import csv
import html as _html
import json
import os
import re
import time

FORMATS = {"pdf", "docx", "xlsx", "pptx", "md", "html", "txt", "csv"}
ALIASES = {"word": "docx", "doc": "docx", "excel": "xlsx", "xls": "xlsx", "spreadsheet": "xlsx", "powerpoint": "pptx",
           "ppt": "pptx", "slides": "pptx", "deck": "pptx", "markdown": "md", "text": "txt", "plain": "txt", "htm": "html"}
MAX_PARTS = 20
MAX_CHARS = 400_000
PART_CALL_WORDS = 550
MAX_CONTINUES = 5

_WRITER_RULES = (
    "You are the writing engine of a document generator. Output ONLY the document text in Markdown - no preface, no "
    "closing remarks, no code fence around the whole thing. Use '#' headings for structure, '-' bullets, numbered lists, "
    "'| a | b |' tables with a header row and a '|---|---|' separator, and ``` fences for code or pseudo-code. "
    "Write real, complete, specific content - never placeholders such as 'lorem ipsum', 'TBD' or 'insert here'. "
    "Honesty rule: never invent citations, quotations, or measured results and present them as real. When the request "
    "calls for numbers or sources you cannot actually know, write them as clearly labelled illustrative or projected "
    "values (say so in the text or table caption), or describe the method instead. Follow every instruction in the "
    "request about tone, structure, length, tables, code blocks and what each section must cover.")

_OUTLINE_SYSTEM = (
    "You plan a document. Reply with ONE JSON object and nothing else: "
    '{"title":"document title","parts":[{"heading":"section heading","brief":"what it must cover, drawn from the request",'
    '"words":300}]}. '
    "One part per top-level section the request asks for, in order (keep the user's own section names and order if "
    "they gave them; include an abstract/introduction/conclusion only if the kind of document has them). "
    f"At most {MAX_PARTS} parts. A short or simple document (a letter, a one-page note, a checklist) is ONE part. "
    "'words' is the target length of that part, sized to the request.")


def norm_format(fmt: str, path: str = "", request: str = "") -> str:
    f = str(fmt or "").lower().strip().lstrip(".")
    f = ALIASES.get(f, f)
    if f in FORMATS:
        return f
    ext = os.path.splitext(str(path or ""))[1].lower().lstrip(".")
    ext = ALIASES.get(ext, ext)
    if ext in FORMATS:
        return ext
    low = (request or "").lower()
    for word, out in (("pdf", "pdf"), ("word", "docx"), (".docx", "docx"), ("excel", "xlsx"), ("spreadsheet", "xlsx"),
                      ("powerpoint", "pptx"), ("slide", "pptx"), ("csv", "csv"), ("markdown", "md"), ("html", "html")):
        if word in low:
            return out
    return ""


def free_name(folder: str, name: str) -> str:
    base, ext = os.path.splitext(name)
    cand, n = os.path.join(folder, name), 1
    while os.path.exists(cand):
        cand = os.path.join(folder, f"{base} ({n}){ext}")
        n += 1
    return cand


def _slug(s: str, n: int = 60) -> str:
    s = str(s or "").replace(":", " -")
    s = re.sub(r'[<>"/\\|?*\x00-\x1f]', "", s).strip(" .")
    s = re.sub(r"\s+", " ", s)[:n].strip(" .")
    return s or "Document"


def strip_fence(text: str) -> str:
    t = (text or "").strip()
    m = re.match(r"^```(?:markdown|md)?\s*\n(.*)\n```\s*$", t, re.S | re.I)
    return m.group(1).strip() if m else t


def extract_json_obj(raw: str):
    t = strip_fence(str(raw or ""))
    a, b = t.find("{"), t.rfind("}")
    if a < 0 or b <= a:
        return None
    try:
        return json.loads(t[a:b + 1])
    except ValueError:
        return None


def ends_clean(md: str) -> bool:
    """False when a reply looks CUT OFF (model hit its token limit): open code fence, or the last line stops mid-sentence."""
    t = (md or "").rstrip()
    if not t:
        return False
    if t.count("```") % 2:
        return False
    last = t.splitlines()[-1].strip()
    if last.startswith(("|", "#", "```")) or re.match(r"^(-{3,}|\*{3,})$", last):
        return True
    if re.match(r"^([-*+]|\d+[.)])\s+", last) and len(last.split()) <= 10:
        return True
    return last[-1] in ".!?:;)]\"'\u201d\u2019*`_"


_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _cells(line: str) -> list:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in re.split(r"(?<!\\)\|", line)]


def parse_blocks(md: str) -> list:
    """-> [("h", level, text) | ("p", text) | ("ul", [items]) | ("ol", [items]) | ("table", [[cells]]) | ("code", text) | ("quote", text) | ("hr",)]"""
    lines = (md or "").replace("\r\n", "\n").replace("\t", "    ").split("\n")
    out, i = [], 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            out.append(("code", "\n".join(buf).rstrip()))
            continue
        m = re.match(r"^(#{1,6})\s+(.*?)\s*#*\s*$", s)
        if m:
            out.append(("h", len(m.group(1)), m.group(2)))
            i += 1
            continue
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
            out.append(("hr",))
            i += 1
            continue
        if "|" in s and i + 1 < len(lines) and _TABLE_SEP.match(lines[i + 1]) and "|" in lines[i + 1]:
            rows = [_cells(s)]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(_cells(lines[i]))
                i += 1
            w = max(len(r) for r in rows)
            out.append(("table", [r + [""] * (w - len(r)) for r in rows]))
            continue
        if re.match(r"^[-*+]\s+", s) or re.match(r"^\d+[.)]\s+", s):
            ordered = bool(re.match(r"^\d+[.)]\s+", s))
            items = []
            while i < len(lines):
                t = lines[i].strip()
                mm = re.match(r"^(?:[-*+]|\d+[.)])\s+(.*)$", t)
                if mm:
                    items.append(mm.group(1))
                elif t and lines[i].startswith("  ") and items:
                    items[-1] += " " + t
                else:
                    break
                i += 1
            out.append(("ol" if ordered else "ul", items))
            continue
        if s.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            out.append(("quote", " ".join(buf)))
            continue
        buf = [s]
        i += 1
        while i < len(lines):
            t = lines[i].strip()
            if (not t or t.startswith(("```", "#", ">")) or re.match(r"^([-*+]|\d+[.)])\s+", t)
                    or ("|" in t and i + 1 < len(lines) and _TABLE_SEP.match(lines[i + 1]))):
                break
            buf.append(t)
            i += 1
        out.append(("p", " ".join(buf)))
    return out


_INLINE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*|_[^_\s][^_]*_|`[^`]+`)")


def inline_runs(text: str) -> list:
    """'a **b** `c`' -> [(text, bold, italic, code)]"""
    runs = []
    for part in _INLINE.split(text or ""):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            runs.append((part[2:-2], True, False, False))
        elif part.startswith("__") and part.endswith("__") and len(part) > 4:
            runs.append((part[2:-2], True, False, False))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            runs.append((part[1:-1], False, False, True))
        elif part[0] in "*_" and part[-1] == part[0] and len(part) > 2:
            runs.append((part[1:-1], False, True, False))
        else:
            runs.append((part, False, False, False))
    return runs


def plain(text: str) -> str:
    return "".join(r[0] for r in inline_runs(text))


def _words(blocks) -> int:
    n = 0
    for b in blocks:
        if b[0] in ("p", "quote", "code"):
            n += len(str(b[-1]).split())
        elif b[0] == "h":
            n += len(b[2].split())
        elif b[0] in ("ul", "ol"):
            n += sum(len(x.split()) for x in b[1])
        elif b[0] == "table":
            n += sum(len(c.split()) for r in b[1] for c in r)
    return n


def _find_font(names, dirs):
    for d in dirs:
        for n in names:
            p = os.path.join(d, n)
            if os.path.isfile(p):
                return p
    return None


def _font_set():
    """A Unicode TTF family (regular/bold/italic/bold-italic) + a monospace one, searched in the usual Windows/Linux/macOS places.
    Returns None if none is found (the PDF then uses built-in Latin-1 fonts and unsupported characters are replaced)."""
    dirs = [os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts"), "/usr/share/fonts/truetype/dejavu",
            "/usr/share/fonts/dejavu", "/Library/Fonts", "/System/Library/Fonts/Supplemental", os.path.expanduser("~/.fonts")]
    fams = [("times.ttf", "timesbd.ttf", "timesi.ttf", "timesbi.ttf"),
            ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf"),
            ("DejaVuSerif.ttf", "DejaVuSerif-Bold.ttf", "DejaVuSerif-Italic.ttf", "DejaVuSerif-BoldItalic.ttf"),
            ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf", "DejaVuSans-BoldOblique.ttf")]
    for fam in fams:
        reg = _find_font([fam[0]], dirs)
        if reg:
            return ([reg] + [_find_font([f], dirs) or reg for f in fam[1:]],
                    _find_font(["consola.ttf", "cour.ttf", "DejaVuSansMono.ttf"], dirs))
    return None


_LATIN = {"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"', "\u2013": "-", "\u2014": "--", "\u2026": "...",
          "\u00a0": " ", "\u2192": "->", "\u2264": "<=", "\u2265": ">=", "\u2022": "-", "\u00d7": "x"}


def _latin(t: str) -> str:
    for k, v in _LATIN.items():
        t = t.replace(k, v)
    return t.encode("latin-1", "ignore").decode("latin-1")


def render_pdf(title: str, blocks: list, path: str):
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos
    fonts = _font_set()
    pdf = FPDF(format="A4")
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(True, 20)
    if fonts:
        fam, mono = fonts
        for style, p in zip(("", "B", "I", "BI"), fam):
            pdf.add_font("Body", style, p)
        if mono:
            pdf.add_font("Mono", "", mono)
        body, mono_name, fix = "Body", ("Mono" if mono else "Body"), (lambda s: s)
    else:
        body, mono_name, fix = "Helvetica", "Courier", _latin
    pdf.set_title(title)
    pdf.add_page()

    def para(text, size=11, h=5.6, indent=0.0, style=""):
        pdf.set_font(body, style, size)
        pdf.set_x(pdf.l_margin + indent)
        md = re.sub(r"`([^`]+)`", r"\1", text)
        md = re.sub(r"(?<![*\w])\*([^*\s][^*]*)\*(?![*\w])", r"__\1__", md)
        try:
            pdf.multi_cell(0, h, fix(md), markdown=True, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        except Exception:
            pdf.set_x(pdf.l_margin + indent)
            pdf.multi_cell(0, h, fix(plain(text)), new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.set_font(body, "B", 20)
    pdf.multi_cell(0, 9, fix(title), align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(5)
    sizes = {1: 15, 2: 13, 3: 12, 4: 11, 5: 11, 6: 11}
    for b in blocks:
        k = b[0]
        if k == "h":
            pdf.ln(3)
            pdf.set_font(body, "B", sizes[b[1]])
            pdf.multi_cell(0, 7, fix(plain(b[2])), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(1.5)
        elif k == "p":
            para(b[1])
            pdf.ln(2)
        elif k in ("ul", "ol"):
            for n, it in enumerate(b[1], 1):
                pdf.set_font(body, "", 11)
                bullet = "-" if k == "ul" else f"{n}."
                pdf.set_x(pdf.l_margin + 3)
                pdf.cell(6, 5.6, bullet)
                y = pdf.get_y()
                pdf.set_xy(pdf.l_margin + 9, y)
                pdf.set_left_margin(pdf.l_margin + 9)
                para(it)
                pdf.set_left_margin(pdf.l_margin - 9)
            pdf.ln(1.5)
        elif k == "quote":
            para(b[1], style="I", indent=6)
            pdf.ln(2)
        elif k == "code":
            pdf.set_font(mono_name, "", 8.5)
            pdf.set_fill_color(242, 242, 242)
            for ln in (b[1] or " ").split("\n"):
                pdf.set_x(pdf.l_margin)
                pdf.multi_cell(0, 4.4, fix(ln.replace("\t", "    ")) or " ", fill=True, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(3)
        elif k == "hr":
            pdf.ln(1)
            pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
            pdf.ln(3)
        elif k == "table":
            rows = b[1]
            pdf.set_font(body, "", 9)
            try:
                with pdf.table(first_row_as_headings=True, text_align="LEFT", line_height=4.8, padding=1) as t:
                    for r in rows:
                        row = t.row()
                        for c in r:
                            row.cell(fix(plain(c)))
            except Exception:
                for r in rows:
                    para(" | ".join(plain(c) for c in r), size=9, h=4.8)
            pdf.ln(3)
    pdf.set_font(body, "", 9)
    pdf.output(path)


def render_docx(title: str, blocks: list, path: str):
    import docx
    from docx.shared import Pt

    def add_runs(par, text, mono=False):
        for t, bold, ital, code in inline_runs(text):
            r = par.add_run(t)
            r.bold, r.italic = bold or None, ital or None
            if code or mono:
                r.font.name = "Consolas"
    d = docx.Document()
    d.core_properties.title = title
    d.add_heading(title, level=0)
    for b in blocks:
        k = b[0]
        if k == "h":
            add_runs(d.add_heading("", level=min(b[1], 9)), b[2])
        elif k == "p":
            add_runs(d.add_paragraph(), b[1])
        elif k in ("ul", "ol"):
            for it in b[1]:
                add_runs(d.add_paragraph(style="List Bullet" if k == "ul" else "List Number"), it)
        elif k == "quote":
            add_runs(d.add_paragraph(style="Intense Quote"), b[1])
        elif k == "code":
            for ln in (b[1] or " ").split("\n"):
                p = d.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(ln)
                r.font.name, r.font.size = "Consolas", Pt(9)
            d.add_paragraph()
        elif k == "hr":
            d.add_paragraph("_" * 40)
        elif k == "table":
            rows = b[1]
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for ri, r in enumerate(rows):
                for ci, c in enumerate(r):
                    cell = t.cell(ri, ci)
                    cell.text = ""
                    add_runs(cell.paragraphs[0], c)
                    if ri == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            d.add_paragraph()
    d.save(path)


def _num(v: str):
    s = str(v).strip().replace(",", "")
    if re.fullmatch(r"-?\d+", s) and not re.match(r"-?0\d", s) and len(s.lstrip("-")) <= 15:
        return int(s)
    if re.fullmatch(r"-?\d*\.\d+", s):
        return float(s)
    return plain(v) if not str(v).startswith("=") else str(v)


def render_xlsx(title: str, blocks: list, path: str):
    """Each table becomes a sheet (named after the heading above it). Text-only content goes on one sheet, a line per row."""
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    heading, used = "", set()

    def sheet_name(h):
        base = re.sub(r"[\\/*?:\[\]]", "", h or f"Sheet{len(used) + 1}")[:28] or "Sheet"
        n, name = 1, base
        while name.lower() in used:
            n += 1
            name = f"{base[:25]}_{n}"
        used.add(name.lower())
        return name
    for b in blocks:
        if b[0] == "h":
            heading = plain(b[2])
        elif b[0] == "table":
            ws = wb.create_sheet(sheet_name(heading))
            for ri, r in enumerate(b[1], 1):
                for ci, c in enumerate(r, 1):
                    cell = ws.cell(row=ri, column=ci, value=plain(c) if ri == 1 else _num(c))
                    if ri == 1:
                        cell.font = Font(bold=True)
            for ci in range(1, len(b[1][0]) + 1):
                width = max(len(str(ws.cell(row=r, column=ci).value or "")) for r in range(1, min(ws.max_row, 200) + 1))
                ws.column_dimensions[get_column_letter(ci)].width = min(max(10, width + 2), 60)
            ws.freeze_panes = "A2"
    if not wb.sheetnames:
        ws = wb.create_sheet(sheet_name(title))
        for b in blocks:
            if b[0] in ("ul", "ol"):
                for it in b[1]:
                    ws.append([plain(it)])
            elif b[0] in ("p", "h", "quote"):
                ws.append([plain(b[-1])])
        ws.column_dimensions["A"].width = 100
    wb.save(path)


def render_pptx(title: str, blocks: list, path: str):
    """Title slide, then one slide per heading (level 1-2) with its bullets/paragraphs; tables become real tables."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = title
    s.shapes.title.left, s.shapes.title.width = Inches(0.7), Inches(11.9)
    if len(s.placeholders) > 1:
        s.placeholders[1].text = ""
        s.placeholders[1].left, s.placeholders[1].width = Inches(0.7), Inches(11.9)
    slides, cur = [], None
    for b in blocks:
        if b[0] == "h" and b[1] <= 2:
            cur = {"title": plain(b[2]), "items": [], "tables": [], "code": []}
            slides.append(cur)
            continue
        if cur is None:
            cur = {"title": title, "items": [], "tables": [], "code": []}
            slides.append(cur)
        if b[0] == "h":
            cur["items"].append(plain(b[2]))
        elif b[0] == "p":
            cur["items"].append(plain(b[1]))
        elif b[0] in ("ul", "ol"):
            cur["items"] += [plain(x) for x in b[1]]
        elif b[0] == "table":
            cur["tables"].append(b[1])
        elif b[0] == "code":
            cur["code"].append(b[1])
    for sl in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = sl["title"][:90]
        slide.shapes.title.left, slide.shapes.title.top = Inches(0.7), Inches(0.4)
        slide.shapes.title.width, slide.shapes.title.height = Inches(11.9), Inches(1.0)
        top = Inches(1.5)
        if sl["items"]:
            box = slide.shapes.add_textbox(Inches(0.7), top, Inches(11.9), Inches(2.6 if sl["tables"] else 5.4))
            tf = box.text_frame
            tf.word_wrap = True
            for n, it in enumerate(sl["items"][:12]):
                p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
                p.text = "\u2022 " + it
                p.font.size = Pt(20 if len(sl["items"]) <= 6 else 16)
            top = Inches(4.2) if sl["tables"] else top
        for rows in sl["tables"][:1]:
            rows = rows[:10]
            tbl = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.7), top, Inches(11.9), Inches(0.4) * len(rows)).table
            for ri, r in enumerate(rows):
                for ci, c in enumerate(r):
                    tbl.cell(ri, ci).text = plain(c)
                    for p in tbl.cell(ri, ci).text_frame.paragraphs:
                        p.font.size = Pt(12)
        if sl["code"] and not sl["tables"]:
            box = slide.shapes.add_textbox(Inches(0.7), Inches(4.4) if sl["items"] else top, Inches(11.9), Inches(2.6))
            box.text_frame.text = sl["code"][0][:900]
            for p in box.text_frame.paragraphs:
                p.font.size, p.font.name = Pt(12), "Consolas"
    prs.save(path)


def render_md(title: str, blocks: list, path: str, raw_md: str = ""):
    body = raw_md.strip()
    if not re.match(r"^#\s", body):
        body = f"# {title}\n\n{body}"
    with open(path, "w", encoding="utf-8") as f:
        f.write(body + "\n")


def render_txt(title: str, blocks: list, path: str):
    out = [title, "=" * len(title), ""]
    for b in blocks:
        k = b[0]
        if k == "h":
            out += ["", plain(b[2]), ("-" if b[1] > 1 else "=") * len(plain(b[2]))]
        elif k in ("p", "quote"):
            out += [plain(b[1]), ""]
        elif k in ("ul", "ol"):
            out += [(f"  {n}. " if k == "ol" else "  - ") + plain(it) for n, it in enumerate(b[1], 1)] + [""]
        elif k == "code":
            out += ["    " + ln for ln in b[1].split("\n")] + [""]
        elif k == "table":
            w = [max(len(plain(r[c])) for r in b[1]) for c in range(len(b[1][0]))]
            out += ["  ".join(plain(c).ljust(w[i]) for i, c in enumerate(r)) for r in b[1]] + [""]
        elif k == "hr":
            out.append("-" * 40)
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(out).rstrip() + "\n")


def _h(text: str) -> str:
    t = _html.escape(text or "")
    t = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"`([^`]+)`", r"<code>\1</code>", t)
    return re.sub(r"(?<![*\w])\*([^*\s][^*]*)\*(?![*\w])", r"<i>\1</i>", t)


def render_html(title: str, blocks: list, path: str):
    parts = []
    for b in blocks:
        k = b[0]
        if k == "h":
            parts.append(f"<h{min(b[1] + 1, 6)}>{_h(b[2])}</h{min(b[1] + 1, 6)}>")
        elif k == "p":
            parts.append(f"<p>{_h(b[1])}</p>")
        elif k in ("ul", "ol"):
            parts.append(f"<{k}>" + "".join(f"<li>{_h(i)}</li>" for i in b[1]) + f"</{k}>")
        elif k == "quote":
            parts.append(f"<blockquote>{_h(b[1])}</blockquote>")
        elif k == "code":
            parts.append(f"<pre><code>{_html.escape(b[1])}</code></pre>")
        elif k == "hr":
            parts.append("<hr>")
        elif k == "table":
            head, *rows = b[1]
            parts.append("<table><thead><tr>" + "".join(f"<th>{_h(c)}</th>" for c in head) + "</tr></thead><tbody>"
                         + "".join("<tr>" + "".join(f"<td>{_h(c)}</td>" for c in r) + "</tr>" for r in rows) + "</tbody></table>")
    css = ("body{font-family:Georgia,serif;max-width:820px;margin:2rem auto;padding:0 1rem;line-height:1.55;color:#222}"
           "pre{background:#f4f4f4;padding:.8rem;overflow:auto}table{border-collapse:collapse;margin:1rem 0}"
           "td,th{border:1px solid #bbb;padding:.35rem .6rem}th{background:#eee}blockquote{border-left:3px solid #ccc;margin-left:0;padding-left:1rem;color:#555}")
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"<!doctype html><html><head><meta charset='utf-8'><title>{_html.escape(title)}</title><style>{css}</style></head>"
                f"<body><h1>{_html.escape(title)}</h1>{''.join(parts)}</body></html>")


def render_csv(title: str, blocks: list, path: str):
    tables = [b[1] for b in blocks if b[0] == "table"]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        if tables:
            for r in tables[0]:
                w.writerow([plain(c) for c in r])
        else:
            for b in blocks:
                if b[0] in ("ul", "ol"):
                    for it in b[1]:
                        w.writerow([plain(it)])
                elif b[0] in ("p", "h"):
                    w.writerow([plain(b[-1])])


RENDERERS = {"pdf": render_pdf, "docx": render_docx, "xlsx": render_xlsx, "pptx": render_pptx, "txt": render_txt,
             "html": render_html, "csv": render_csv}
NEEDS_BINARY = {"pdf", "docx", "xlsx", "pptx"}
NEEDS = {"pdf": "fpdf2", "docx": "python-docx", "xlsx": "openpyxl", "pptx": "python-pptx"}


class DocWriter:
    def __init__(self, llm_fn, default_dir: str, safe_path=None, objective_fn=None, progress_fn=None, on_created=None, log=print):
        """llm_fn(system, user) -> text.  safe_path(path) -> real path or raises (the registry's sandbox).
        objective_fn() -> the user's FULL original request (so a planner that shortened it can't lose requirements)."""
        self.llm, self.default_dir, self.safe_path = llm_fn, os.path.abspath(default_dir), safe_path
        self.objective_fn, self.progress, self.log = objective_fn, progress_fn or (lambda m: None), log
        self.on_created = on_created

    def _ask(self, system: str, user: str) -> str:
        err = None
        for _ in range(2):
            try:
                out = str(self.llm(system, user) or "").strip()
                if out:
                    return out
                err = "empty reply"
            except Exception as e:
                err = f"{type(e).__name__}: {str(e)[:100]}"
                time.sleep(1.5)
        raise RuntimeError(err or "the model did not answer")

    def _brief(self, request: str) -> str:
        full = ""
        try:
            full = str(self.objective_fn() or "").strip() if self.objective_fn else ""
        except Exception:
            full = ""
        request = (request or "").strip()
        if full and len(full) > len(request) * 1.15 and request[:40].lower() not in full.lower():
            return f"THE USER'S FULL REQUEST (authoritative):\n{full}\n\nPLANNER'S SUMMARY OF THE JOB:\n{request}"
        if full and len(full) > len(request):
            return f"THE USER'S FULL REQUEST (authoritative):\n{full}"
        return request

    def _outline(self, brief: str, kind: str) -> dict:
        for extra in ("", " Reply with the JSON object ONLY - no prose, no code fence."):
            try:
                obj = extract_json_obj(self._ask(_OUTLINE_SYSTEM + f" The output is a {kind}." + extra, brief))
            except RuntimeError:
                obj = None
            if isinstance(obj, dict):
                return obj
        return {}

    def _chunked(self, system: str, first_user: str, words: int, base_user: str) -> str:
        """One section, however long. A call only ever asks for PART_CALL_WORDS; when the section needs more - or the reply was
        cut off mid-sentence - it CONTINUES from the tail instead of silently handing back half a section."""
        md = strip_fence(self._ask(system, first_user))
        for _ in range(MAX_CONTINUES):
            have, clean = len(md.split()), ends_clean(md)
            left = words - have
            if len(md) > MAX_CHARS or (clean and left < words * 0.4):
                break
            need = max(80, min(left, PART_CALL_WORDS))
            user = (base_user + f"\n\nWHAT YOU HAVE WRITTEN SO FAR ENDS WITH:\n{md[-700:]}\n\n"
                    + ("Your last sentence was cut off: finish it, then " if not clean else "")
                    + f"CONTINUE exactly where it stops with about {need} more words. Do not repeat anything and do not restart the "
                      "section. If this part is already complete, reply with exactly: DONE")
            more = strip_fence(self._ask(system, user))
            if more.strip().upper().rstrip(".") == "DONE":
                break
            md = md + (" " if not clean else "\n\n") + more.lstrip()
        return md

    def write(self, request: str, fmt: str, title: str = "") -> tuple:
        """-> (title, markdown)"""
        brief = self._brief(request)
        if fmt in ("xlsx", "csv"):
            sysm = (_WRITER_RULES + " The output is a spreadsheet: give each sheet a '## Sheet name' heading followed by ONE markdown "
                    "table (header row first, one record per row, real values, numbers without units or thousands separators; "
                    "formulas may be written as =SUM(B2:B9)). Add no prose between sheets.")
            return title or "Workbook", strip_fence(self._ask(sysm, brief))
        kind = "slide deck: '## ' heading = one slide, 3-6 short bullets or one table per slide" if fmt == "pptx" else "document"
        outline = self._outline(brief, kind)
        parts = [p for p in (outline.get("parts") or []) if isinstance(p, dict) and str(p.get("heading", "")).strip()][:MAX_PARTS]
        title = _slug(title or outline.get("title") or "", 100) if (title or outline.get("title")) else ""
        if len(parts) <= 1 or fmt == "pptx":
            try:
                words = int((parts[0] if parts else {}).get("words") or 500)
            except (TypeError, ValueError):
                words = 500
            sysm = _WRITER_RULES + f" The output is a {kind}. Do not repeat the title as a heading; start with the first section."
            user = brief + (f"\n\nTitle: {title}" if title else "")
            md = self._chunked(sysm, user, 0 if fmt == "pptx" else words, user)
            return title or self._title_from(md, request), md
        plan = "\n".join(f"{n}. {p['heading']} - {p.get('brief', '')}" for n, p in enumerate(parts, 1))
        chunks, tail, total = [], "", 0
        for n, p in enumerate(parts, 1):
            self.progress(f"Writing section {n} of {len(parts)}: {p['heading']}")
            try:
                words = int(p.get("words") or 300)
            except (TypeError, ValueError):
                words = 300
            words = max(80, min(words, 1500))
            base = (f"{brief}\n\nDOCUMENT TITLE: {title or '(untitled)'}\nFULL OUTLINE:\n{plan}\n\n"
                    + (f"END OF THE PREVIOUS SECTION (continue naturally, do not repeat it):\n{tail}\n\n" if tail else "")
                    + f"YOU ARE WRITING ONLY section {n}: '{p['heading']}'. It must cover: {p.get('brief', '')}\n"
                      f"The whole section is about {words} words. Begin with the line '# {p['heading']}' "
                      "(use '##' and '###' for sub-headings), then the content. Do not write other sections.")
            first = base + f"\nWrite about {min(words, PART_CALL_WORDS)} words now" + (" (the section continues afterwards)." if words > PART_CALL_WORDS else ".")
            try:
                md = self._chunked(_WRITER_RULES, first, words, base)
            except RuntimeError as e:
                raise RuntimeError(f"the model failed while writing section {n} of {len(parts)} ('{p['heading']}'): {e}") from None
            if not re.match(r"^#{1,6}\s", md):
                md = f"# {p['heading']}\n\n{md}"
            chunks.append(md)
            tail = md[-700:]
            total += len(md)
            if total > MAX_CHARS:
                break
        return title or "Document", "\n\n".join(chunks)

    @staticmethod
    def _title_from(md: str, request: str) -> str:
        m = re.search(r"^#\s+(.+)$", md, re.M)
        if m:
            return _slug(plain(m.group(1)), 100)
        return _slug(re.sub(r"\s+", " ", request)[:60], 60)

    def _resolve(self, path: str, fmt: str, title: str) -> tuple:
        """-> (folder, file name). Pure: touches nothing on disk."""
        p = os.path.expanduser(str(path or "").strip())
        folder, name = self.default_dir, ""
        if p:
            full = p if os.path.isabs(p) else os.path.join(self.default_dir, p)
            is_dir = os.path.isdir(full) or full.endswith(("\\", "/"))
            if is_dir or not os.path.splitext(full)[1]:
                folder = full.rstrip("\\/") if is_dir else os.path.dirname(full)
                name = "" if is_dir else os.path.basename(full)
            else:
                folder, name = os.path.dirname(full), os.path.basename(full)
        return folder, _slug(os.path.splitext(name)[0] or title) + "." + fmt

    def check(self, path: str, fmt: str) -> str | None:
        """Is this place allowed? Asked BEFORE the model writes anything, so a protected folder costs no model calls."""
        try:
            folder, name = self._resolve(path, fmt or "txt", "Document")
            if self.safe_path:
                self.safe_path(os.path.join(folder, name))
        except (PermissionError, ValueError, OSError) as e:
            return f"{type(e).__name__}: {e}"
        return None

    def target(self, path: str, fmt: str, title: str) -> str:
        folder, name = self._resolve(path, fmt, title)
        if self.safe_path:
            self.safe_path(os.path.join(folder, name))
        os.makedirs(folder, exist_ok=True)
        return free_name(folder, name)

    def create(self, request: str = "", format: str = "", path: str = "", title: str = "", content: str = "") -> dict:
        fmt = norm_format(format, path, request)
        if not fmt:
            return {"ok": False, "error": f"I need a format: one of {', '.join(sorted(FORMATS))}"}
        if not (str(content).strip() or str(request).strip()):
            return {"ok": False, "error": "tell me what the document should contain"}
        bad = self.check(path, fmt)
        if bad:
            return {"ok": False, "error": bad}
        miss = NEEDS.get(fmt)
        if miss:
            mod = {"fpdf2": "fpdf", "python-docx": "docx", "openpyxl": "openpyxl", "python-pptx": "pptx"}[miss]
            try:
                __import__(mod)
            except ImportError:
                return {"ok": False, "error": f"a library is missing for {fmt} files - run: pip install {miss}"}
        if str(content).strip():
            md, doc_title = strip_fence(str(content)), _slug(title) if title else ""
            doc_title = doc_title or self._title_from(md, request)
        else:
            doc_title, md = self.write(str(request), fmt, str(title or ""))
        blocks = parse_blocks(md)
        if not blocks:
            return {"ok": False, "error": "the content came out empty"}
        if blocks[0][0] == "h" and blocks[0][1] == 1 and plain(blocks[0][2]).strip().lower() == doc_title.strip().lower():
            blocks = blocks[1:]
        out = self.target(path, fmt, doc_title)
        tmp = out + ".part"
        try:
            if fmt == "md":
                render_md(doc_title, blocks, tmp, md)
            else:
                RENDERERS[fmt](doc_title, blocks, tmp)
            if not os.path.isfile(tmp) or os.path.getsize(tmp) < (50 if fmt in NEEDS_BINARY else 1):
                raise RuntimeError("the file came out empty")
            os.replace(tmp, out)
        except Exception as e:
            try:
                os.path.exists(tmp) and os.remove(tmp)
            except OSError:
                pass
            return {"ok": False, "error": f"couldn't build the {fmt} file: {type(e).__name__}: {str(e)[:160]}"}
        info = {"path": out, "name": os.path.basename(out), "format": fmt, "title": doc_title, "words": _words(blocks),
                "sections": sum(1 for b in blocks if b[0] == "h" and b[1] <= 2), "bytes": os.path.getsize(out)}
        if fmt == "pdf":
            try:
                from pypdf import PdfReader
                info["pages"] = len(PdfReader(out).pages)
            except Exception:
                pass
        info["text"] = (f"Saved {fmt.upper()} '{info['name']}' to {out} - {info['words']} words, {info['sections']} sections"
                        + (f", {info['pages']} pages" if info.get("pages") else ""))
        if self.on_created:
            try:
                self.on_created(dict(info))
            except Exception as e:
                self.log(f"[DOCS] on_created hook failed: {type(e).__name__}: {str(e)[:100]}")
        return {"ok": True, "data": info}


def register(reg, *, llm_fn, default_dir: str, objective_fn=None, progress_fn=None, on_created=None, log=print) -> DocWriter:
    import xenora_tools as xt
    w = DocWriter(llm_fn, default_dir, safe_path=getattr(reg, "safe_path", None), objective_fn=objective_fn,
                  progress_fn=progress_fn, on_created=on_created, log=log)

    def _precheck(a):
        fmt = norm_format(a.get("format"), a.get("path"), a.get("request"))
        if not fmt:
            return f"'format' must be one of: {', '.join(sorted(FORMATS))}"
        if not (str(a.get("content") or "").strip() or str(a.get("request") or "").strip()):
            return "'request' must say what the document should contain"
        return w.check(a.get("path"), fmt)

    reg.register(xt.Tool(
        "create_document",
        "Create a NEW document file (pdf, docx/Word, xlsx/Excel, pptx/PowerPoint, md, html, txt, csv) and SAVE it. ONE step does everything: "
        "it writes the content itself, however long (reports, papers, essays, letters, spreadsheets, slide decks), then builds the file. "
        "Put the user's whole request in `request` (all requirements, section list, tone, tables). Do NOT write the content yourself, "
        "and do NOT search or read first unless the user asked you to use a source - then pass that text as `content` instead "
        "(`content` is saved as-is, no writing). Never overwrites (a name that exists becomes 'name (1)'); returns the saved path "
        "(data.path) - use that path in later steps. Long documents take a few minutes. Not for editing a file that exists.",
        lambda request="", format="", path="", title="", content="": w.create(request, format, path, title, content),
        {"request": "everything the user wants in the document", "format": "pdf | docx | xlsx | pptx | md | html | txt | csv",
         "path": "folder or file path to save to (optional; default = my documents folder)", "title": "document title (optional)",
         "content": "ready-made text/Markdown to save as-is instead of writing it (optional)"},
        ["format"], risk="act", precheck=_precheck,
        approval=lambda a: {"domain": "files", "op": "create",
                            "label": f"create {norm_format(a.get('format'), a.get('path'), a.get('request')) or 'a'} document",
                            "value": "", "desc": "Create that document file?"}))
    return w
